import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

# Load the JSON data
with open('weather_forecast.json', 'r') as jsonfile:
    data = json.load(jsonfile)

# Create a DataFrame from the JSON data
df = pd.DataFrame(data)

# Create a new Excel workbook
workbook = Workbook()
worksheet = workbook.active

# Convert the DataFrame to rows
rows = dataframe_to_rows(df, index=False, header=True)

# Write the rows to the worksheet
for r_idx, row in enumerate(rows, 1):
    for c_idx, value in enumerate(row, 1):
        worksheet.cell(row=r_idx, column=c_idx, value=value)

# Create a table in the worksheet
table = Table(displayName="WeatherForecast", ref=worksheet.dimensions)
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2")
worksheet.add_table(table)

# Save the workbook to an Excel file
workbook.save('weather_forecast.xlsx')

print("Excel file generated successfully.")
